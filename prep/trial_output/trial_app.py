"""MONALISA Paris 2025 series browser (Streamlit).

Read-only browser for data/MONALISA_Paris_2025.csv. No cleaning, no gap
filling, no unit conversion. Notes are stored in notes.json next to this file.
"""

from __future__ import annotations

import io
import json
import os
import re
import tempfile
from datetime import datetime

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import streamlit as st
from docx import Document
from docx.shared import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
CSV_PATH = os.path.join(HERE, "data", "MONALISA_Paris_2025.csv")
NOTES_PATH = os.path.join(HERE, "notes.json")

GROUPS = {
    "NH4_": "NH4+ reagent ions",
    "H3O_": "H3O+ reagent ions",
    "gas_": "Trace gases",
    "met_": "Meteorology",
    "n_": "Sample counters",
}
TZ_NOTE = "Timestamps as stored in the file; the file states no time zone."
WIND_DIR = "met_Wind_Direction (degrees)"
ACCUMULATED = {
    "met_Precipitation (mm)",
    "met_Sunshine_Duration (min)",
    "met_Global_Radiation (J/cm2)",
}
RESOLUTIONS = {"raw (30 min)": None, "hourly": "1h", "daily": "1D"}


# --------------------------------------------------------------------------- #
# data
# --------------------------------------------------------------------------- #
@st.cache_data(show_spinner="Reading CSV ...")
def load_data() -> pd.DataFrame:
    df = pd.read_csv(CSV_PATH)
    df["time"] = pd.to_datetime(df["time"])
    return df.set_index("time").sort_index()


def group_of(col: str) -> str:
    for prefix in GROUPS:
        if col.startswith(prefix):
            return prefix
    return "other"


def unit_of(col: str):
    m = re.search(r"\(([^()]*)\)\s*$", col)
    return m.group(1) if m else None


def axis_label(col: str) -> str:
    unit = unit_of(col)
    if unit:
        return f"{col} [{unit}]"
    if col.startswith("n_"):
        return f"{col} [count of raw samples]"
    return f"{col} [unit not stated in file]"


def circular_mean(values) -> float:
    """Vector mean of compass bearings in degrees, returned in (0, 360]."""
    v = pd.Series(values).dropna().to_numpy(dtype=float)
    if v.size == 0:
        return float("nan")
    rad = np.deg2rad(v)
    ang = float(np.rad2deg(np.arctan2(np.sin(rad).mean(), np.cos(rad).mean())) % 360.0)
    return 360.0 if ang == 0.0 else ang


def aggregate(series: pd.Series, rule):
    """Resample one series.

    Plain arithmetic mean of the non-null values in each bin, with two
    exceptions: wind direction uses a circular (vector) mean, and the n_*
    sample counters are summed because they are counts, not measurements.
    """
    if rule is None:
        return series
    name = str(series.name)
    if name == WIND_DIR:
        return series.resample(rule).apply(circular_mean)
    if name.startswith("n_"):
        return series.resample(rule).sum(min_count=1)
    return series.resample(rule).mean()


def diurnal(series: pd.Series) -> pd.DataFrame:
    """Mean value per hour of day, weekdays and weekends separately.

    The hour of day is the hour of the timestamp exactly as stored in the file
    (no time-zone conversion is applied or possible). Aggregation rules are the
    same as in `aggregate`.
    """
    name = str(series.name)
    s = series.dropna()
    out = pd.DataFrame(index=pd.Index(range(24), name="hour"))
    for label, mask in (
        ("weekdays (Mon-Fri)", s.index.dayofweek < 5),
        ("weekend (Sat-Sun)", s.index.dayofweek >= 5),
    ):
        part = s[mask]
        if part.empty:
            out[label] = np.nan
            out[label + " n"] = 0
            continue
        grouped = part.groupby(part.index.hour)
        if name == WIND_DIR:
            vals = grouped.apply(circular_mean)
        elif name.startswith("n_"):
            vals = grouped.sum()
        else:
            vals = grouped.mean()
        out[label] = vals.reindex(out.index)
        out[label + " n"] = grouped.size().reindex(out.index).fillna(0).astype(int)
    return out


def stats_of(series: pd.Series, n_rows: int) -> dict:
    s = series.dropna()
    if s.empty:
        return {"n": 0, "coverage %": 0.0}
    return {
        "n": int(s.size),
        "coverage %": round(100.0 * s.size / n_rows, 2),
        "mean": float(s.mean()),
        "median": float(s.median()),
        "std": float(s.std()),
        "min": float(s.min()),
        "max": float(s.max()),
        "first valid": str(s.index[0]),
        "last valid": str(s.index[-1]),
    }


# --------------------------------------------------------------------------- #
# notes
# --------------------------------------------------------------------------- #
def load_notes() -> dict:
    if not os.path.exists(NOTES_PATH):
        return {}
    try:
        with open(NOTES_PATH, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def save_note(series_name: str, author: str, text: str) -> None:
    """Append a note and rewrite notes.json atomically (tmp file + replace)."""
    notes = load_notes()
    notes.setdefault(series_name, []).append(
        {
            "author": author,
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "text": text,
        }
    )
    fd, tmp = tempfile.mkstemp(dir=HERE, suffix=".json")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            json.dump(notes, fh, indent=2, ensure_ascii=False)
            fh.flush()
            os.fsync(fh.fileno())
        os.replace(tmp, NOTES_PATH)
    except BaseException:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise


# --------------------------------------------------------------------------- #
# figures
# --------------------------------------------------------------------------- #
def timeseries_figure(data: pd.DataFrame, cols, rule, res_label: str):
    fig, axes = plt.subplots(
        len(cols), 1, figsize=(9, 2.4 * len(cols)), sharex=True, squeeze=False
    )
    for ax, col in zip(axes[:, 0], cols):
        s = aggregate(data[col], rule)
        ax.plot(s.index, s.to_numpy(dtype=float), lw=0.8)
        ax.set_ylabel(axis_label(col), fontsize=7)
        ax.grid(alpha=0.3)
    axes[-1, 0].set_xlabel(f"time -- {res_label}. {TZ_NOTE}", fontsize=8)
    fig.tight_layout()
    return fig


def diurnal_figure(data: pd.DataFrame, cols):
    fig, axes = plt.subplots(
        len(cols), 1, figsize=(9, 2.4 * len(cols)), sharex=True, squeeze=False
    )
    for ax, col in zip(axes[:, 0], cols):
        d = diurnal(data[col])
        for label in ("weekdays (Mon-Fri)", "weekend (Sat-Sun)"):
            ax.plot(d.index, d[label].to_numpy(dtype=float), marker="o", ms=3, lw=1, label=label)
        ax.set_ylabel(axis_label(col), fontsize=7)
        ax.grid(alpha=0.3)
        ax.legend(fontsize=6)
    axes[-1, 0].set_xticks(range(0, 24, 3))
    axes[-1, 0].set_xlabel(f"hour of day. {TZ_NOTE}", fontsize=8)
    fig.tight_layout()
    return fig


def fig_png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=130)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# report
# --------------------------------------------------------------------------- #
def build_report(data: pd.DataFrame, cols, rule, res_label: str, notes: dict, figures) -> bytes:
    doc = Document()
    doc.add_heading("MONALISA Paris 2025 -- series notes", level=0)
    doc.add_paragraph(f"Generated {datetime.now().isoformat(timespec='seconds')}")

    for col in cols:
        doc.add_heading(col, level=1)
        doc.add_paragraph(f"Unit: {unit_of(col) or 'not stated in file'}")
        doc.add_heading("Notes", level=2)
        entries = notes.get(col, [])
        if entries:
            for e in entries:
                doc.add_paragraph(
                    f"[{e['timestamp']}] {e['author']}: {e['text']}", style="List Bullet"
                )
        else:
            doc.add_paragraph("(no notes)")
        doc.add_heading("Statistics", level=2)
        for k, v in stats_of(data[col], len(data)).items():
            doc.add_paragraph(f"{k}: {v}", style="List Bullet")

    doc.add_heading("Figures on screen", level=1)
    for png in figures:
        doc.add_picture(io.BytesIO(png), width=Inches(6.2))

    doc.add_heading("Provenance", level=1)
    for line in [
        f"Source file: {CSV_PATH}",
        f"Rows: {len(data)}; columns: {len(data.columns)}",
        f"Period covered: {data.index[0]} to {data.index[-1]} (native grid 30 min)",
        TZ_NOTE,
        f"Resolution plotted: {res_label}",
        "Processing: no cleaning, gap filling, outlier removal or unit conversion; "
        "values are used as read from the CSV.",
        "Aggregation: arithmetic mean of the non-null values per bin, except "
        f"'{WIND_DIR}', which uses a circular (vector) mean, and the n_* counters, "
        "which are summed.",
        "Diurnal cycle: mean per hour of day of the timestamp as stored, weekdays "
        "(Mon-Fri) and weekend (Sat-Sun) computed separately.",
        "Note: the met_ columns are hourly values repeated on both half-hour slots; "
        "met_Precipitation, met_Sunshine_Duration and met_Global_Radiation are hourly "
        "accumulations and are averaged, not summed.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# app
# --------------------------------------------------------------------------- #
st.set_page_config(page_title="MONALISA browser", layout="wide")
df = load_data()
all_cols = list(df.columns)

st.sidebar.header("Series")
chosen_groups = st.sidebar.multiselect(
    "Groups",
    list(GROUPS),
    default=["gas_", "met_"],
    format_func=lambda g: f"{g} -- {GROUPS[g]}",
)
name_filter = st.sidebar.text_input("Name contains")
candidates = [
    c
    for c in all_cols
    if (not chosen_groups or group_of(c) in chosen_groups)
    and name_filter.lower() in c.lower()
]
st.sidebar.caption(f"{len(candidates)} of {len(all_cols)} series match")
selected = st.sidebar.multiselect("Selected series (max 4)", candidates, max_selections=4)

st.title("MONALISA Paris 2025 -- series browser")
st.caption(
    f"{os.path.basename(CSV_PATH)} | {len(df)} rows | "
    f"{df.index[0]} to {df.index[-1]} | {TZ_NOTE}"
)

tab_series, tab_scatter, tab_notes = st.tabs(
    ["Time series & diurnal", "Scatter", "Notes"]
)

with tab_series:
    res_label = st.radio("Resolution", list(RESOLUTIONS), horizontal=True)
    rule = RESOLUTIONS[res_label]
    ts_fig = None
    dc_fig = None
    if not selected:
        st.info("Pick one to four series in the sidebar.")
    else:
        ts_fig = timeseries_figure(df, selected, rule, res_label)
        dc_fig = diurnal_figure(df, selected)
        left, right = st.columns(2)
        with left:
            st.subheader("Time series")
            st.pyplot(ts_fig)
        with right:
            st.subheader("Mean diurnal cycle")
            st.pyplot(dc_fig)
        if WIND_DIR in selected:
            st.caption(f"'{WIND_DIR}' is averaged as a circular (vector) mean.")
        if ACCUMULATED.intersection(selected):
            st.caption(
                "Accumulated met variables (precipitation, sunshine duration, global "
                "radiation) are hourly totals in the file and are averaged, not summed."
            )

        st.subheader("Notes for the selected series")
        notes = load_notes()
        for col in selected:
            st.markdown(f"**{col}**")
            entries = notes.get(col, [])
            if not entries:
                st.write("(no notes)")
            for e in entries:
                st.write(f"- [{e['timestamp']}] **{e['author']}**: {e['text']}")

    st.divider()
    if selected and st.button("Export Word report"):
        figures = [fig_png(f) for f in (ts_fig, dc_fig) if f is not None]
        payload = build_report(df, selected, rule, res_label, load_notes(), figures)
        st.download_button(
            "Download .docx",
            data=payload,
            file_name=f"MONALISA_notes_{datetime.now():%Y%m%d_%H%M%S}.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

with tab_scatter:
    if not candidates:
        st.info("No series match the sidebar filter.")
    else:
        c1, c2 = st.columns(2)
        x_col = c1.selectbox("X series", candidates, index=0, key="xcol")
        y_col = c2.selectbox(
            "Y series", candidates, index=1 if len(candidates) > 1 else 0, key="ycol"
        )
        if x_col == y_col:
            st.warning("Pick two different series.")
        else:
            pair = df[[x_col, y_col]].dropna()
            if len(pair) < 2:
                st.warning("Fewer than two timestamps where both series are present.")
            else:
                r = float(pair[x_col].corr(pair[y_col]))
                fig, ax = plt.subplots(figsize=(5.5, 5))
                ax.plot(
                    pair[x_col].to_numpy(dtype=float),
                    pair[y_col].to_numpy(dtype=float),
                    "o",
                    ms=3,
                    alpha=0.4,
                )
                ax.set_xlabel(axis_label(x_col), fontsize=8)
                ax.set_ylabel(axis_label(y_col), fontsize=8)
                ax.grid(alpha=0.3)
                fig.tight_layout()
                st.pyplot(fig)
                st.write(
                    f"Pearson r = {r:.3f}, n = {len(pair)} "
                    "(raw 30-min timestamps where both series are present)"
                )

with tab_notes:
    if not selected:
        st.info("Select series in the sidebar to attach a note.")
    else:
        target = st.selectbox("Series", selected, key="notetarget")
        author = st.text_input("Who said it", key="noteauthor")
        text = st.text_area("Note", key="notetext")
        if st.button("Save note"):
            if not author.strip() or not text.strip():
                st.error("Both a note and an author are required.")
            else:
                save_note(target, author.strip(), text.strip())
                st.success(f"Saved to {os.path.basename(NOTES_PATH)}")
        st.subheader(f"Notes on record for {target}")
        for e in load_notes().get(target, []):
            st.write(f"- [{e['timestamp']}] **{e['author']}**: {e['text']}")
