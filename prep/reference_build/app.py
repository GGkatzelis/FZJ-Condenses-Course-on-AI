"""MONALISA — interactive timeseries browser, built live during the lecture.

Run with:   .venv\\Scripts\\python.exe -m streamlit run app.py

Layout follows Georgios' spec:
  left  : list of available series
  right : two plots - the timeseries, and the weekend/weekday diurnal
  plus  : a correlations subpanel, per-series notes saved to notes.json,
          and a Word overview built from those notes.

CIRCULAR_WIND is the wind-direction trap switch. False reproduces what a live
build writes by default (a scalar mean of a circular variable, wrong by up to
160 degrees). The demo flips it to True once someone in the room notices.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import platform
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import streamlit as st

CIRCULAR_WIND = False

HERE = Path(__file__).parent
CSV = HERE / "data" / "MONALISA_Paris_2025.csv"
NOTES = HERE / "notes.json"
PROMPTS = HERE / "prompts.json"
REPORT = HERE / "MONALISA_overview.docx"

NAVY, ORANGE, GREEN, GREY = "#1F4E79", "#C0622E", "#55733D", "#8A9099"

st.set_page_config(page_title="MONALISA browser", layout="wide")


# ------------------------------------------------------------------ data
@st.cache_data
def load() -> pd.DataFrame:
    d = pd.read_csv(CSV)
    d["time"] = pd.to_datetime(d["time"])
    return d


df = load()
T = "time"
df["_hour"] = df[T].dt.hour
df["_weekend"] = df[T].dt.weekday >= 5

SERIES = [c for c in df.columns if c not in (T, "_hour", "_weekend")]
FAMILY = {"H3O_": "VOCs, H3O+ mode", "NH4_": "VOCs, NH4+ mode",
          "gas_": "Gases and aerosol", "met_": "Meteorology", "n_": "Bookkeeping"}


def family(col: str) -> str:
    for k, v in FAMILY.items():
        if col.startswith(k):
            return v
    return "Other"


def unit_of(col: str) -> str:
    """Units are embedded in the gas_/met_ column names. The ion columns have none."""
    if "(" in col and col.rstrip().endswith(")"):
        return col[col.index("(") + 1: col.rindex(")")]
    return ""


def short(col: str) -> str:
    return col.split(" (")[0]


def is_direction(col: str) -> bool:
    return "Direction" in col


# ----------------------------------------------------------------- notes
def read_json(path: Path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def save_note(series: str, author: str, text: str) -> None:
    """Append a note under its series key and write immediately, so an app
    restart cannot lose it. notes.json is a dict of series -> list of notes."""
    notes = read_json(NOTES, {})
    notes.setdefault(series, []).append({
        "timestamp": dt.datetime.now().isoformat(timespec="seconds"),
        "author": author.strip() or "anon",
        "text": text.strip(),
    })
    tmp = NOTES.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(notes, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(NOTES)


# ----------------------------------------------------------------- stats
def circular_mean(deg: pd.Series) -> float:
    a = np.deg2rad(deg.dropna().to_numpy(dtype=float))
    if a.size == 0:
        return float("nan")
    return float(np.rad2deg(np.arctan2(np.sin(a).mean(), np.cos(a).mean())) % 360)


def diurnal(col: str, weekend: bool) -> pd.Series:
    sub = df[df["_weekend"] == weekend]
    g = sub.groupby("_hour")[col]
    if CIRCULAR_WIND and is_direction(col):
        return g.apply(circular_mean)
    return g.mean()


def n_days(weekend: bool) -> int:
    return df.loc[df["_weekend"] == weekend, T].dt.date.nunique()


# ----------------------------------------------------------------- plots
def tidy(ax, xlabel, ylabel):
    ax.set_xlabel(xlabel)
    ax.set_ylabel(ylabel)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)


def fig_timeseries(cols, resample):
    fig, ax = plt.subplots(figsize=(11, 3.6), dpi=120)
    for c, colour in zip(cols, (NAVY, ORANGE, GREEN, "#7B5AA6")):
        s = df.set_index(T)[c]
        if resample != "30 min":
            s = s.resample({"hourly": "1h", "daily": "1D"}[resample]).mean()
        ax.plot(s.index, s.values, lw=1.0, color=colour, label=short(c))
    units = sorted({unit_of(c) or "no unit" for c in cols})
    tidy(ax, "time", " / ".join(units))
    if len(units) > 1:
        ax.set_title("note: these series are not in the same unit", fontsize=9,
                     color=ORANGE, loc="left")
    ax.legend(frameon=False, fontsize=9, ncols=min(len(cols), 4))
    fig.tight_layout()
    return fig


def fig_diurnal(cols):
    fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=120)
    for c, colour in zip(cols, (NAVY, ORANGE, GREEN, "#7B5AA6")):
        wd, we = diurnal(c, False), diurnal(c, True)
        ax.plot(wd.index, wd.values, lw=2.2, color=colour, marker="o", ms=4,
                label=f"{short(c)} — weekday")
        ax.plot(we.index, we.values, lw=2.2, color=colour, ls="--", marker="s",
                ms=4, alpha=.75, label=f"{short(c)} — weekend")
    tidy(ax, "hour of day", " / ".join(sorted({unit_of(c) or "no unit" for c in cols})))
    ax.set_xticks(range(0, 24, 3))
    ax.legend(frameon=False, fontsize=8)
    fig.tight_layout()
    return fig


def fig_scatter(x, y):
    fig, ax = plt.subplots(figsize=(6.0, 4.8), dpi=120)
    s = df[[x, y, "_hour"]].dropna()
    sc = ax.scatter(s[x], s[y], c=s["_hour"], cmap="twilight_shifted", s=20,
                    lw=0, alpha=.85)
    fig.colorbar(sc, ax=ax).set_label("hour of day")
    r = s[x].corr(s[y])
    rs = s[x].corr(s[y], method="spearman")
    ax.set_title(f"Pearson r = {r:.2f}   Spearman = {rs:.2f}   n = {len(s)}",
                 fontsize=10)
    tidy(ax, f"{short(x)} [{unit_of(x) or 'no unit'}]",
         f"{short(y)} [{unit_of(y) or 'no unit'}]")
    fig.tight_layout()
    return fig, r, rs, len(s)


def png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    return buf.getvalue()


# -------------------------------------------------------------------- UI
st.sidebar.title("Series")
st.sidebar.caption(f"{len(SERIES):,} series in the file — filter, then pick up to 4")

fams = list(dict.fromkeys(family(c) for c in SERIES))
fam_sel = st.sidebar.multiselect("family", fams, default=["Meteorology"],
                                 key="fam")
query = st.sidebar.text_input("name contains", placeholder="e.g. C7H9, NO2, Wind",
                              key="q")

candidates = [c for c in SERIES
              if (not fam_sel or family(c) in fam_sel)
              and (not query or query.lower() in c.lower())]
st.sidebar.caption(f"{len(candidates):,} match")

MAX_LIST = 300
if len(candidates) > MAX_LIST:
    st.sidebar.info(f"Too many to list. Showing the first {MAX_LIST} — "
                    f"narrow the filter to see the rest.")
picked = st.sidebar.multiselect("plot these", candidates[:MAX_LIST],
                                format_func=short, max_selections=4, key="pick")

st.title("MONALISA — timeseries browser")

tab_ts, tab_corr, tab_report = st.tabs(
    ["Time series", "Correlations", "Overview report"])

# ---- Time series + diurnal ----
with tab_ts:
    if not picked:
        st.info("Choose one or more series in the list on the left.")
    else:
        res = st.radio("averaging", ["30 min", "hourly", "daily"],
                       horizontal=True, key="res")
        f1 = fig_timeseries(picked, res)
        st.pyplot(f1)
        st.session_state["fig_ts"] = png(f1)

        left, right = st.columns([1.05, 1])
        with left:
            st.subheader("Weekend vs weekday")
            f2 = fig_diurnal(picked)
            st.pyplot(f2)
            st.session_state["fig_di"] = png(f2)
            st.caption(
                f"Weekday curves average {n_days(False)} days; weekend curves "
                f"average only {n_days(True)}. Dashed = weekend.")
        with right:
            st.subheader("Note on a series")
            with st.form("note_ts", clear_on_submit=True):
                which = st.selectbox("about which series", picked, format_func=short)
                who = st.text_input("who said it", placeholder="name, or 'the room'")
                txt = st.text_area("their interpretation, in their words", height=120)
                if st.form_submit_button("Save note") and txt.strip():
                    save_note(which, who, txt)
                    st.success(f"saved against {short(which)}")

        notes = read_json(NOTES, {})
        for c in picked:
            for nt in notes.get(c, []):
                st.markdown(f"**{short(c)}** — {nt['text']}  \n*{nt['author']}, "
                            f"{nt['timestamp'][11:16]}*")

# ---- Correlations ----
@st.cache_data(show_spinner=False)
def top_correlations(target: str, k: int = 20) -> pd.DataFrame:
    """Rank every other series by |r| against `target`. 880 columns takes ~0.2 s."""
    num = df[SERIES]
    r = num.corrwith(num[target]).drop(index=target, errors="ignore").dropna()
    out = r.reindex(r.abs().sort_values(ascending=False).index).head(k)
    return pd.DataFrame({
        "series": [short(c) for c in out.index],
        "r": out.values.round(3),
        "R²": (out.values ** 2).round(3),
        "_col": out.index,
    })


def fig_ranking(tbl: pd.DataFrame, target: str, hi: int | None = None):
    fig, ax = plt.subplots(figsize=(6.2, 5.6), dpi=120)
    ypos = np.arange(len(tbl))[::-1]
    cols = [ORANGE if (hi is not None and i == hi) else NAVY
            for i in range(len(tbl))]
    ax.barh(ypos, tbl["r"], color=cols, height=.78)
    ax.set_yticks(ypos)
    ax.set_yticklabels(tbl["series"], fontsize=8)
    ax.axvline(0, color="#40454A", lw=.8)
    ax.set_xlabel("Pearson r")
    ax.set_title(f"strongest correlations with {short(target)}", fontsize=11)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    return fig


def fig_fit(x: str, y: str):
    """Scatter with an OLS fit and the statistics that go with it."""
    s = df[[x, y, "_hour"]].dropna()
    fig, ax = plt.subplots(figsize=(6.4, 5.2), dpi=120)
    sc = ax.scatter(s[x], s[y], c=s["_hour"], cmap="twilight_shifted", s=20,
                    lw=0, alpha=.85)
    fig.colorbar(sc, ax=ax).set_label("hour of day")
    stats = {}
    if len(s) > 2:
        slope, intercept = np.polyfit(s[x], s[y], 1)
        xf = np.linspace(s[x].min(), s[x].max(), 50)
        ax.plot(xf, slope * xf + intercept, color=ORANGE, lw=2.4)
        r = s[x].corr(s[y])
        stats = {"slope": slope, "intercept": intercept, "r": r, "R2": r ** 2,
                 "spearman": s[x].corr(s[y], method="spearman"), "n": len(s)}
        ax.set_title(f"slope = {slope:.3g}   R² = {r ** 2:.3f}   n = {len(s)}",
                     fontsize=11)
    tidy(ax, f"{short(x)} [{unit_of(x) or 'no unit'}]",
         f"{short(y)} [{unit_of(y) or 'no unit'}]")
    fig.tight_layout()
    return fig, stats


with tab_corr:
    st.subheader("What correlates with what")
    # Deliberately NOT restricted to the sidebar filter: on stage you want to
    # type a compound straight in without setting the family filter first.
    # Streamlit's selectbox is type-to-search, so 880 entries is fine.
    default = SERIES.index(picked[0]) if picked else 0
    target = st.selectbox("compound  (type to search all series)", SERIES,
                          index=default, format_func=short, key="ctarget")
    tbl = top_correlations(target)

    left, right = st.columns([1, 1.05])
    with left:
        sel = st.dataframe(
            tbl[["series", "r", "R²"]], width="stretch", hide_index=True,
            on_select="rerun", selection_mode="single-row", key="cpick")
        rows = sel.selection.rows if hasattr(sel, "selection") else []
        hi = rows[0] if rows else 0
        st.pyplot(fig_ranking(tbl, target, hi))
        st.caption("Select a row to see the fit on the right.")

    partner = tbl["_col"].iloc[hi]
    with right:
        f3, stats = fig_fit(target, partner)
        st.pyplot(f3)
        st.session_state["fig_sc"] = png(f3)
        if stats:
            st.session_state["corr"] = {"x": target, "y": partner,
                                        "pearson": stats["r"],
                                        "spearman": stats["spearman"],
                                        "n": stats["n"], "slope": stats["slope"],
                                        "R2": stats["R2"]}
            st.write(
                f"**{short(target)}** vs **{short(partner)}**  \n"
                f"slope **{stats['slope']:.4g}**  ·  intercept "
                f"**{stats['intercept']:.4g}**  \n"
                f"Pearson r **{stats['r']:.3f}**  ·  R² **{stats['R2']:.3f}**  \n"
                f"Spearman **{stats['spearman']:.3f}**  ·  n **{stats['n']}**")
            if abs(stats["r"] - stats["spearman"]) > 0.15:
                st.warning(f"Pearson and Spearman differ by "
                           f"{abs(stats['r'] - stats['spearman']):.2f}: not linear, "
                           f"or a few extreme points dominate.")
        st.caption("These are half-hourly samples and strongly autocorrelated, "
                   "so n overstates the degrees of freedom. No p-value is shown "
                   "for that reason.")

    with st.form("note_corr", clear_on_submit=True):
        who = st.text_input("who said it", key="cwho")
        txt = st.text_area("interpretation of this correlation", height=90)
        if st.form_submit_button("Save note") and txt.strip():
            save_note(f"{short(target)} vs {short(partner)}", who, txt)
            st.success("saved")

# ---- Word overview ----
with tab_report:
    st.subheader("Word overview of our notes")
    notes = read_json(NOTES, {})
    st.write(f"**{sum(len(v) for v in notes.values())} note(s)** across "
             f"**{len(notes)} series**")
    st.json(notes, expanded=False)

    def build() -> Path:
        from docx import Document
        from docx.shared import Inches, Pt

        doc = Document()
        doc.add_heading("MONALISA — what we found", 0)
        doc.add_paragraph(
            f"Compiled live during “AI tools for scientists”, "
            f"{dt.date.today().isoformat()}.")

        doc.add_heading("Figures", 1)
        for key, cap in (("fig_ts", "Time series"),
                         ("fig_di", "Weekend vs weekday diurnal"),
                         ("fig_sc", "Correlation")):
            if st.session_state.get(key):
                doc.add_heading(cap, 2)
                doc.add_picture(io.BytesIO(st.session_state[key]), width=Inches(6.2))

        doc.add_heading("Notes, by series", 1)
        if not notes:
            doc.add_paragraph("No notes were recorded.")
        for series, items in notes.items():
            doc.add_heading(short(series), 2)
            if series in df.columns:
                s = df[series].dropna()
                stat = (f"n = {len(s)}, mean = {s.mean():.3g}, "
                        f"median = {s.median():.3g}, max = {s.max():.3g}")
                if is_direction(series):
                    stat += (f", vector mean = {circular_mean(s):.0f} deg"
                             if CIRCULAR_WIND else "")
                doc.add_paragraph(f"{stat}  [{unit_of(series) or 'no unit given'}]"
                                  ).runs[0].font.size = Pt(9)
            for nt in items:
                doc.add_paragraph(nt["text"], style="Intense Quote")
                doc.add_paragraph(f"— {nt['author']}, {nt['timestamp']}"
                                  ).runs[0].font.size = Pt(9)

        cr = st.session_state.get("corr")
        if cr:
            doc.add_heading("Correlation examined", 1)
            doc.add_paragraph(
                f"{short(cr['x'])} vs {short(cr['y'])}: Pearson r = "
                f"{cr['pearson']:.3f}, Spearman = {cr['spearman']:.3f}, "
                f"n = {cr['n']}.")

        doc.add_heading("Provenance", 1)
        rows = [
            ("Data file", CSV.name),
            ("SHA-256", hashlib.sha256(CSV.read_bytes()).hexdigest()[:32] + "…"),
            ("Rows × columns", f"{len(df)} × {len(SERIES) + 1}"),
            ("Date range", f"{df[T].min()} to {df[T].max()}"),
            ("Time base", "30 minutes"),
            ("Time zone", "not stated in the data file"),
            ("Averaging shown", st.session_state.get("res", "30 min")),
            ("Weekend definition", "Saturday and Sunday by the timestamp as given"),
            ("Days averaged", f"{n_days(False)} weekday, {n_days(True)} weekend"),
            ("Wind direction averaging",
             "vector (circular) mean" if CIRCULAR_WIND
             else "arithmetic mean of degrees"),
            ("Missing data", "excluded pairwise; gaps are not filled"),
            ("Tool", "Claude Code in VS Code"),
            ("Model", "Claude Opus 5"),
            ("Python / pandas", f"{platform.python_version()} / {pd.__version__}"),
            ("Generated", dt.datetime.now().isoformat(timespec="seconds")),
        ]
        tb = doc.add_table(rows=0, cols=2)
        tb.style = "Light List Accent 1"
        for k, v in rows:
            cells = tb.add_row().cells
            cells[0].text, cells[1].text = k, str(v)

        doc.add_heading("Prompts used", 2)
        prompts = read_json(PROMPTS, [])
        for i, p in enumerate(prompts, 1):
            doc.add_paragraph(f"{i}. {p}")
        if not prompts:
            doc.add_paragraph("prompts.json not found — prompts were not recorded.")

        doc.save(REPORT)
        return REPORT

    if st.button("Build the Word overview", type="primary"):
        p = build()
        st.success(f"written: {p.name}")
        st.download_button("download it", p.read_bytes(), file_name=p.name,
                           mime="application/vnd.openxmlformats-officedocument."
                                "wordprocessingml.document")
